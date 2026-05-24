"""
scripts/Smoke tests/smoke_file_analysis.py
Smoke test manuale per modules/file_analysis.

Test in due modalità:

1) **modalità batch (default)** — genera file di esempio in una cartella
   temporanea e prova ogni estrattore non-audio (PDF, DOCX, TXT, MD,
   HTML, JSON, CSV, XML). Niente network, niente Whisper.

2) **modalità file** — passa uno o più path reali da analizzare. Se i
   path includono file audio (.wav/.mp3/...) e c'è --with-stt, WhisperSTT
   viene caricato e usato per trascrivere.

Esecuzione:
    cd ~/assistant

    # Batch su file generati: estrattori non-audio
    venv-runtime/bin/python "scripts/Smoke tests/smoke_file_analysis.py"

    # Su uno o più file reali
    venv-runtime/bin/python "scripts/Smoke tests/smoke_file_analysis.py" \\
        ~/Scaricati/contratto.pdf ~/Documenti/note.md

    # Con STT (per file audio)
    venv-runtime/bin/python "scripts/Smoke tests/smoke_file_analysis.py" \\
        --with-stt ~/Musica/registrazione.mp3

    # Mostra contenuto esteso (default: solo i primi 400 char per file)
    venv-runtime/bin/python "scripts/Smoke tests/smoke_file_analysis.py" \\
        --full ~/Scaricati/contratto.pdf
"""

from __future__ import annotations

import argparse
import asyncio
import json
import sys
import tempfile
import time
from pathlib import Path

# Setup path del progetto
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from modules.file_analysis import AnalysisResult, FileAnalysisError, FileAnalyzer


# ---------------------------------------------------------------------------
# Helpers di stampa
# ---------------------------------------------------------------------------

W = 72  # larghezza header


def header(title: str) -> None:
    print()
    print("=" * W)
    print(f"  {title}")
    print("=" * W)


def section(title: str) -> None:
    print()
    print(f"--- {title} ".ljust(W, "-"))


def ok(msg: str) -> None:
    print(f"  ✓ {msg}")


def warn(msg: str) -> None:
    print(f"  ⚠  {msg}")


def fail(msg: str) -> None:
    print(f"  ✗ {msg}")


def kv(k: str, v: object) -> None:
    print(f"    {k:<20} {v}")


def show_result(r: AnalysisResult, *, full: bool = False, preview_chars: int = 400) -> None:
    """Stampa un AnalysisResult in modo umano-leggibile."""
    if r.error:
        fail(f"errore: {r.error}")
        kv("path",  r.path)
        kv("tipo",  r.file_type)
        return

    ok(f"{r.file_type:<5}  {r.path}")
    kv("estrattore",   r.extractor)
    kv("caratteri",    f"{r.char_count} (originale: {r.original_char_count})")
    kv("troncato",     "sì" if r.truncated else "no")
    kv("tempo",        f"{r.elapsed_ms:.1f} ms")
    kv("file size",    f"{r.metadata.get('file_size_bytes', '?')} byte")

    print()
    if full or len(r.content) <= preview_chars:
        print("    ── contenuto ──")
        for line in r.content.splitlines():
            print(f"    {line}")
    else:
        print(f"    ── contenuto (primi {preview_chars} char di {len(r.content)}) ──")
        snippet = r.content[:preview_chars]
        for line in snippet.splitlines():
            print(f"    {line}")
        print(f"    [...]")


# ---------------------------------------------------------------------------
# Generatori di file di esempio per la modalità batch
# ---------------------------------------------------------------------------

def make_sample_files(d: Path) -> dict[str, Path]:
    """
    Genera un set rappresentativo di file in `d`.
    Ritorna dict { etichetta: path }.
    """
    samples: dict[str, Path] = {}

    # TXT
    (d / "note.txt").write_text(
        "Riga uno: appunti dello smoke.\n"
        "Riga due: testo unicode → mercoledì, caffè, città.\n",
        encoding="utf-8",
    )
    samples["TXT"] = d / "note.txt"

    # MD
    (d / "readme.md").write_text(
        "# Smoke test\n\n"
        "Questo è un file Markdown di prova per il modulo file_analysis.\n"
        "- punto uno\n- punto due\n",
        encoding="utf-8",
    )
    samples["MD"] = d / "readme.md"

    # HTML
    (d / "page.html").write_text(
        "<html><head><style>body{color:red}</style></head>"
        "<body><h1>Titolo</h1><p>Paragrafo visibile.</p>"
        "<script>alert(1)</script></body></html>",
        encoding="utf-8",
    )
    samples["HTML"] = d / "page.html"

    # JSON
    (d / "data.json").write_text(
        json.dumps({"nome": "Mauro", "ruolo": "dev", "lingua": "italiano"}),
        encoding="utf-8",
    )
    samples["JSON"] = d / "data.json"

    # CSV
    (d / "people.csv").write_text(
        "nome,ruolo,anni\n"
        "Mauro,dev,42\n"
        "Gwen,assistente,30\n",
        encoding="utf-8",
    )
    samples["CSV"] = d / "people.csv"

    # XML
    (d / "config.xml").write_text(
        "<config><user>mauro</user><lang>it</lang></config>",
        encoding="utf-8",
    )
    samples["XML"] = d / "config.xml"

    # PDF (vuoto, solo per verificare il wrapper pypdf)
    try:
        from pypdf import PdfWriter
        w = PdfWriter()
        w.add_blank_page(width=72, height=72)
        with open(d / "blank.pdf", "wb") as fh:
            w.write(fh)
        samples["PDF"] = d / "blank.pdf"
    except Exception as exc:
        warn(f"PDF: skip generazione ({exc})")

    # DOCX
    try:
        from docx import Document
        doc = Document()
        doc.add_paragraph("Paragrafo di prova per lo smoke test.")
        doc.add_paragraph("Seconda riga con caratteri italiani: perché, città.")
        doc.save(str(d / "doc.docx"))
        samples["DOCX"] = d / "doc.docx"
    except Exception as exc:
        warn(f"DOCX: skip generazione ({exc})")

    return samples


# ---------------------------------------------------------------------------
# Sezioni di test
# ---------------------------------------------------------------------------

async def test_metadata(fa: FileAnalyzer) -> None:
    header("1. metadata del FileAnalyzer")
    kv("max_chars_inline", fa._max_chars)
    kv("max_file_bytes",   fa._max_file_bytes)
    kv("safe_dirs",        fa._safe_dirs)
    kv("stt iniettato",    "sì" if fa._stt is not None else "no")
    kv("estensioni",       f"{len(fa.supported_extensions())} totali")
    print()
    print("    Estensioni gestite:")
    exts = sorted(fa.supported_extensions())
    line = "      "
    for ext in exts:
        if len(line) + len(ext) + 2 > 60:
            print(line)
            line = "      "
        line += ext + "  "
    if line.strip():
        print(line)


async def test_batch(fa: FileAnalyzer, *, full: bool) -> None:
    """Genera file di esempio in temp dir e li analizza tutti."""
    header("2. modalità batch — file generati al volo")

    with tempfile.TemporaryDirectory(prefix="smoke_file_analysis_") as tmp:
        tmp_path = Path(tmp)
        section(f"generazione file in {tmp_path}")
        samples = make_sample_files(tmp_path)
        ok(f"creati {len(samples)} file: {', '.join(samples.keys())}")

        # Per ogni file: analyze() e stampa
        for label, path in samples.items():
            section(f"{label}  →  {path.name}")
            try:
                r = await fa.analyze(path)
            except Exception as exc:
                fail(f"analyze() ha sollevato (NON dovrebbe!): {exc}")
                continue
            show_result(r, full=full)


async def test_error_paths(fa: FileAnalyzer) -> None:
    """Verifica i path d'errore senza far crashare nulla."""
    header("3. percorsi d'errore (devono essere gestiti)")

    # 3a — file inesistente
    section("3a — file inesistente")
    r = await fa.analyze("/tmp/__non_esiste__.pdf")
    kv("error", r.error)
    if r.error and "not_found" in r.error:
        ok("not_found gestito correttamente")
    else:
        fail("attesi error='[not_found] ...', error attuale: " + str(r.error))

    # 3b — estensione non supportata
    section("3b — estensione non supportata")
    with tempfile.NamedTemporaryFile(suffix=".xyz", delete=False) as fh:
        fh.write(b"ciao")
        bad_path = fh.name
    try:
        r = await fa.analyze(bad_path)
        kv("error", r.error)
        if r.error and "unsupported" in r.error:
            ok("unsupported gestito correttamente")
        else:
            fail("atteso '[unsupported] ...'")
    finally:
        Path(bad_path).unlink(missing_ok=True)

    # 3c — file troppo grande (limite ad hoc)
    section("3c — file troppo grande (max_file_bytes=2)")
    fa_strict = FileAnalyzer(safe_dirs=[], max_file_bytes=2)
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as fh:
        fh.write(b"abcdefghij")
        big_path = fh.name
    try:
        r = await fa_strict.analyze(big_path)
        kv("error", r.error)
        if r.error and "too_large" in r.error:
            ok("too_large gestito correttamente")
        else:
            fail("atteso '[too_large] ...'")
    finally:
        Path(big_path).unlink(missing_ok=True)
        await fa_strict.aclose()

    # 3d — out of safe_dirs
    section("3d — fuori da safe_dirs")
    fa_locked = FileAnalyzer(safe_dirs=["/usr"])
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as fh:
        fh.write(b"ok")
        outside_path = fh.name
    try:
        r = await fa_locked.analyze(outside_path)
        kv("error", r.error)
        if r.error and "out_of_safe_dirs" in r.error:
            ok("out_of_safe_dirs gestito correttamente")
        else:
            fail("atteso '[out_of_safe_dirs] ...'")
    finally:
        Path(outside_path).unlink(missing_ok=True)
        await fa_locked.aclose()


async def test_user_paths(
    fa: FileAnalyzer,
    paths: list[str],
    *,
    full: bool,
) -> None:
    """Analizza i path passati dall'utente."""
    header(f"2. analisi di {len(paths)} file passati dall'utente")

    for path in paths:
        section(path)
        t0 = time.perf_counter()
        try:
            r = await fa.analyze(path)
        except Exception as exc:
            fail(f"analyze() ha sollevato (NON dovrebbe!): {exc}")
            continue
        dt = (time.perf_counter() - t0) * 1000
        show_result(r, full=full)
        print()
        kv("elapsed totale (esterno)", f"{dt:.1f} ms")


# ---------------------------------------------------------------------------
# Setup STT (opzionale)
# ---------------------------------------------------------------------------

async def maybe_load_stt(enable: bool):
    """Carica WhisperSTT se richiesto. Ritorna None se non richiesto/fallisce."""
    if not enable:
        return None

    print("Caricamento WhisperSTT… (può richiedere 5-15s)")
    try:
        from modules.stt import WhisperSTT
        stt = WhisperSTT()
        await stt._load_models()
        ok("WhisperSTT caricato")
        return stt
    except Exception as exc:
        warn(f"impossibile caricare STT: {exc}")
        return None


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

async def main(args: argparse.Namespace) -> int:
    header("smoke_file_analysis.py — modules/file_analysis")

    # STT opzionale per audio
    stt = await maybe_load_stt(args.with_stt)

    # FileAnalyzer con safe_dirs=[] in modalità test (consente tutto, niente
    # interferenze con settings.pc_control.safe_dirs)
    fa = FileAnalyzer(stt=stt, safe_dirs=[])

    t_start = time.monotonic()
    try:
        await test_metadata(fa)

        if args.paths:
            await test_user_paths(fa, args.paths, full=args.full)
        else:
            await test_batch(fa, full=args.full)

        await test_error_paths(fa)
    finally:
        await fa.aclose()
        if stt is not None:
            try:
                await stt.__aexit__(None, None, None)
            except Exception:
                pass

    print()
    print("=" * W)
    elapsed = time.monotonic() - t_start
    ok(f"smoke completato in {elapsed:.1f}s")
    print("=" * W)
    return 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Smoke test del modulo file_analysis",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Esempi:
  %(prog)s
      Modalità batch: genera file di esempio e prova ogni estrattore.

  %(prog)s ~/Scaricati/contratto.pdf
      Analizza un singolo file reale.

  %(prog)s --with-stt ~/Musica/x.mp3
      Carica WhisperSTT e trascrive un file audio.

  %(prog)s --full ~/Documenti/lungo.pdf
      Mostra l'intero contenuto estratto invece dell'anteprima.
        """,
    )
    parser.add_argument(
        "paths",
        nargs="*",
        help="Uno o più path da analizzare. Se omesso, modalità batch.",
    )
    parser.add_argument(
        "--with-stt",
        action="store_true",
        help="Carica WhisperSTT (richiesto per file audio)",
    )
    parser.add_argument(
        "--full",
        action="store_true",
        help="Stampa l'intero contenuto estratto (default: primi 400 char)",
    )
    a = parser.parse_args()

    try:
        sys.exit(asyncio.run(main(a)))
    except KeyboardInterrupt:
        print("\n  Interrotto dall'utente.")
        sys.exit(0)
    except Exception as exc:
        print(f"\n  ❌ Errore fatale: {exc}")
        raise
